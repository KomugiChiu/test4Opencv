#!/usr/bin/env python3
"""Render report_cpp.json (auto_cpp suite) into report_cpp.xlsx.

Usage:
    python3 make_report_xlsx.py --json DIR/report_cpp.json [--out DIR/report_cpp.xlsx]

Sheets:
    Summary  - meta (device/backend/opencv version+source) + counters
    Results  - group/api/status/message per check, color-coded
"""
import argparse
import datetime
import json
import os
import socket


def load_workbook_cls():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        return Workbook, Alignment, Border, Font, PatternFill, Side
    except ImportError:
        import subprocess
        import sys
        subprocess.run([sys.executable, "-m", "pip", "install", "--user",
                        "-q", "openpyxl"])
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (Alignment, Border, Font,
                                         PatternFill, Side)
            return Workbook, Alignment, Border, Font, PatternFill, Side
        except ImportError:
            return None


def render(json_path, out_path):
    cls = load_workbook_cls()
    if cls is None:
        return None
    Workbook, Alignment, Border, Font, PatternFill, Side = cls

    d = json.load(open(json_path, encoding="utf-8"))
    meta = d.get("meta", {})
    summ = d.get("summary", {})
    results = d.get("results", [])

    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"),
             "FAIL": PatternFill("solid", fgColor="FFC7CE"),
             "WARN": PatternFill("solid", fgColor="FFEB9C"),
             "SKIP": PatternFill("solid", fgColor="D9D9D9")}

    counts = {k: summ.get(k, 0) for k in ("PASS", "FAIL", "WARN", "SKIP")}
    total_checks = summ.get("total_checks", len(results))
    executed = summ.get("executed", counts["PASS"] + counts["FAIL"])
    core_total = summ.get("core_total", "")
    ratio = 100.0 * counts["PASS"] / executed if executed else 0.0
    cov = 100.0 * executed / core_total if core_total else 0.0

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["OpenCV Camera API Test Report (C++)", ""])
    ws.append(["timestamp", datetime.datetime.now().isoformat(timespec="seconds")])
    ws.append(["host", socket.gethostname()])
    ws.append(["device", meta.get("device", "")])
    ws.append(["backend", meta.get("backend", "")])
    ws.append(["frames", meta.get("frames", "")])
    ws.append(["opencv version", meta.get("opencv_version", "")])
    ws.append(["opencv source", meta.get("opencv_source", "")])
    ws.append(["", ""])
    for k in ("PASS", "FAIL", "WARN", "SKIP"):
        ws.append([k.lower(), counts[k]])
    ws.append(["total checks", total_checks])
    ws.append(["core coverage", f"{cov:.1f}% ({executed}/{core_total})"])
    ws.append(["passed ratio (excl. skip/warn)", f"{ratio:.1f}%"])
    ws["A1"].font = Font(bold=True, size=14)
    for i in range(2, ws.max_row + 1):
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.cell(row=i, column=1).border = thin
        ws.cell(row=i, column=2).border = thin
    for i, w in enumerate([30, 60], 1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("Results")
    ws2.append(["Group", "API", "Status", "Message"])
    for r in results:
        st = r.get("status", "?")
        if st not in fills:
            fills[st] = PatternFill("solid", fgColor="E7E6E6")
        ws2.append([r.get("group", ""), r.get("api", ""), st,
                    r.get("message", "")])
    for c in ws2[1]:
        c.fill, c.font, c.border = hdr_fill, hdr_font, thin
    for row in ws2.iter_rows(min_row=2, max_col=4):
        st = row[2].value
        if st in fills:
            row[2].fill = fills[st]
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(cell.column == 4))
    for i, w in enumerate([10, 46, 9, 90], 1):
        from openpyxl.utils import get_column_letter
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    wb.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--json", required=True, help="path to report_cpp.json")
    ap.add_argument("--out", default=None,
                    help="output xlsx (default: alongside json)")
    args = ap.parse_args()
    out = args.out or os.path.splitext(args.json)[0] + ".xlsx"
    res = render(args.json, out)
    if res:
        print(res)
        return 0
    print("(openpyxl unavailable; skipped xlsx)", file=sys.stderr)
    return 3


if __name__ == "__main__":
    import sys
    sys.exit(main())
