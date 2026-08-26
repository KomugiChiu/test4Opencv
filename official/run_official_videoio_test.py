#!/usr/bin/env python3
"""Run OpenCV official videoio tests and emit an Excel report.

Usage:
    python3 run_official_videoio_test.py [-e ENV_DIR] [-f GTEST_FILTER]
                                         [-d DEVICE] [-o OUTDIR] [-x EXTRA]

Exit code mirrors the test binary: 0 = all pass, 1 = has failures.
"""
import argparse
import datetime
import os
import re
import socket
import subprocess
import sys

OK_RE = re.compile(r"^\[\s+OK\s+\] (\S+) \((\d+) ms\)")
FAIL_RE = re.compile(r"^\[\s+FAILED\s+\] (\S+?)(?:, .*)? \((\d+) ms\)")
RUN_RE = re.compile(r"^\[\s+RUN\s+\] (\S+)")
# gtest banner printed by modules/ts (SystemInfoCollector::OnTestProgramStart)
VER_RE = re.compile(r"^OpenCV version:\s*(.+)$", re.M)
VCS_RE = re.compile(r"^OpenCV VCS version:\s*(.+)$", re.M)
BUILD_TYPE_RE = re.compile(r"^Build type:\s*(.+)$", re.M)
COMPILER_RE = re.compile(r"^Compiler:\s*(.+)$", re.M)


def opencv_meta_from_log(log_text):
    out = []
    for rx in (VER_RE, VCS_RE, BUILD_TYPE_RE, COMPILER_RE):
        m = rx.search(log_text)
        out.append(m.group(1).strip() if m else "")
    return out  # version, vcs, build_type, compiler


def opencv_version_from_build(build_root):
    """Fallback: read version from the build tree's OpenCVConfig.cmake."""
    path = os.path.join(build_root, "OpenCVConfig.cmake")
    try:
        with open(path, encoding="utf-8", errors="replace") as f:
            m = re.search(r'set\(OpenCV_VERSION\s+"?([0-9][^\s"\)]*)', f.read())
        return m.group(1) if m else ""
    except OSError:
        return ""


def parse_args():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("-e", "--env", default=os.getcwd(),
                    help="env root containing opencv_extra/testdata "
                         "(default: current directory)")
    ap.add_argument("-b", "--build-dir", default=None,
                    help="OpenCV build dir containing bin/opencv_test_videoio "
                         "(default: <env>/build; override for selfbuild versioned paths)")
    ap.add_argument("-f", "--filter", default=None,
                    help="gtest filter, e.g. '*videoio_v4l2*'")
    ap.add_argument("-d", "--device", default=None,
                    help="camera path or index N (auto-converted to /dev/videoN); "
                         "sets OPENCV_TEST_V4L2_VIVID_DEVICE (e.g. /dev/video2 or 2)")
    ap.add_argument("-o", "--outdir", default="./report/new",
                    help="report output directory (default: ./report/new)")
    ap.add_argument("-x", "--extra", default=None,
                    help='extra args passed verbatim to the test binary')
    ap.add_argument("--console-output", action="store_true",
                    help="stream gtest output live to console (default: silent, log only)")
    return ap.parse_args()


def resolve_device(raw):
    text = str(raw).strip()
    if re.fullmatch(r"\d+", text):
        return f"/dev/video{text}"
    return text


def run_tests(binary, data_path, args):
    env = dict(os.environ, OPENCV_TEST_DATA_PATH=data_path)
    if args.device:
        env["OPENCV_TEST_V4L2_VIVID_DEVICE"] = resolve_device(args.device)
    cmd = [binary]
    if args.filter:
        cmd.append(f"--gtest_filter={args.filter}")
    if args.extra:
        cmd.extend(args.extra.split())
    resolved = resolve_device(args.device) if args.device else None
    dev_display = resolved if resolved else "<unset>"
    if args.device and resolved and resolved != args.device:
        dev_display += f" (raw={args.device})"
    print("== binary :", binary)
    print("== data   :", data_path)
    print("== device :", dev_display)
    print("== filter :", args.filter or "<all>")
    live = getattr(args, "console_output", False)
    proc = subprocess.Popen(cmd, env=env,
                            stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT,
                            text=True, errors="replace", bufsize=1)
    log_lines = []
    for line in proc.stdout:
        log_lines.append(line)
        if live:
            sys.stdout.write(line)
            sys.stdout.flush()
    proc.wait()
    return "".join(log_lines), proc.returncode


def summarize(log_text):
    counts, duration, disabled = {}, "?", "0"
    for pat, key in ((r"^\[==========\] (\d+) tests from (\d+) test cases? ran",
                      None),):
        m = re.search(pat, log_text, re.M)
        if m:
            counts["ran"] = m.group(1)
    m = re.search(r"\((\d+) ms total\)", log_text)
    duration = m.group(1) if m else "?"
    for pat, key in ((r"^\[  PASSED  \] (\d+) tests", "passed"),
                     (r"^\[  FAILED  \] (\d+) tests", "failed")):
        m = re.search(pat, log_text, re.M)
        counts[key] = m.group(1) if m else "0"
    m = re.search(r"YOU HAVE (\d+) DISABLED TESTS", log_text)
    disabled = m.group(1) if m else "0"
    return counts, duration, disabled


def collect_cases(lines):
    cases, err_buf, cur = [], [], None
    for ln in lines:
        mo = OK_RE.match(ln)
        mf = FAIL_RE.match(ln)
        if mo or mf:
            status = "PASS" if mo else "FAIL"
            rx = mo or mf
            cases.append({"name": rx.group(1), "status": status,
                          "ms": int(rx.group(2))})
            if status == "FAIL":
                detail = "\n".join(err_buf[-8:]) or "(no captured detail)"
                cur_fail = rx.group(1)
                for c in cases:
                    if c["name"] == cur_fail and c["status"] == "FAIL" \
                            and not c.get("detail"):
                        c["detail"] = detail
            err_buf, cur = [], None
            continue
        mr = RUN_RE.match(ln)
        if mr:
            cur = mr.group(1)
            err_buf = []
        elif cur is not None:
            err_buf.append(ln)
    return cases


def write_excel(path, meta_rows, cases):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "--user",
                        "-q", "openpyxl"])
        try:
            from openpyxl import Workbook
            from openpyxl.styles import (Alignment, Border, Font,
                                         PatternFill, Side)
        except ImportError:
            return None

    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"),
             "FAIL": PatternFill("solid", fgColor="FFC7CE")}

    def header(ws):
        for c in ws[1]:
            c.fill, c.font, c.border = hdr_fill, hdr_font, thin

    def widths(ws, ws_widths):
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for r in meta_rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    for i in range(2, len(meta_rows) + 1):
        ws.cell(row=i, column=1).font = Font(bold=True)
        ws.cell(row=i, column=1).border = thin
        ws.cell(row=i, column=2).border = thin
    widths(ws, [22, 90])

    ws2 = wb.create_sheet("Failed Tests")
    ws2.append(["Test name", "Error detail (last lines of RUN block)"])
    for c in sorted((c for c in cases if c["status"] == "FAIL"),
                    key=lambda x: x["name"]):
        ws2.append([c["name"], c.get("detail", "(see log)")])
    header(ws2)
    for row in ws2.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="9C0006")
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(cell.column == 2))
    widths(ws2, [55, 100])

    ws3 = wb.create_sheet("All Tests")
    ws3.append(["Test name", "Status", "Duration (ms)"])
    for c in sorted(cases, key=lambda x: (-x["ms"], x["name"])):
        ws3.append([c["name"], c["status"], c["ms"]])
    header(ws3)
    for row in ws3.iter_rows(min_row=2, max_col=3):
        st = row[1].value
        if st in fills:
            row[1].fill = fills[st]
        for cell in row:
            cell.border = thin
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = ws3.dimensions
    widths(ws3, [60, 9, 14])

    wb.save(path)
    return path


def main():
    args = parse_args()
    env_root = os.path.abspath(args.env)
    # Build dir: explicit -b overrides, else default to <env>/build
    if args.build_dir:
        build_root = os.path.abspath(args.build_dir)
    else:
        build_root = os.path.join(env_root, "build")
    binary = os.path.join(build_root, "bin", "opencv_test_videoio")
    data_path = os.path.join(env_root, "opencv_extra", "testdata")
    if not os.path.isfile(binary):
        sys.exit(f"ERROR: {binary} not found. "
                 f"Run setup_official_videoio_env.py or build opencv first.")
    if not os.path.isdir(data_path):
        sys.exit(f"ERROR: {data_path} not found.")

    os.makedirs(args.outdir, exist_ok=True)
    log_path = os.path.join(args.outdir, "videoio_gtest.log")
    log_text, rc = run_tests(binary, data_path, args)
    with open(log_path, "w", encoding="utf-8", errors="replace") as f:
        f.write(log_text)

    counts, duration, disabled = summarize(log_text)
    cases = collect_cases(log_text.splitlines())
    ocv_ver, ocv_vcs, ocv_btype, ocv_comp = opencv_meta_from_log(log_text)
    if not ocv_ver:
        ocv_ver = opencv_version_from_build(build_root)
    print(f"""
== result ==
opencv={ocv_ver or '?'}{' (' + ocv_btype + ')' if ocv_btype else ''}
ran={counts.get('ran', '0')} passed={counts['passed']} failed={counts['failed']} \
disabled={disabled} duration_ms={duration} exit={rc}""")

    meta_rows = [
        ["OpenCV Official videoio Test Report", ""],
        ["timestamp", datetime.datetime.now().isoformat(timespec="seconds")],
        ["host", socket.gethostname()],
        ["binary", binary],
        ["opencv version", ocv_ver or "(unknown)"],
        ["opencv vcs", ocv_vcs or "-"],
        ["build type", ocv_btype or "-"],
        ["compiler", ocv_comp or "-"],
        ["gtest_filter", args.filter or "(all)"],
        ["vivid device", resolve_device(args.device) if args.device else "(unset)"],
        ["testdata", data_path],
        ["", ""],
        ["tests ran", counts.get("ran", "0")],
        ["passed", counts["passed"]],
        ["failed", counts["failed"]],
        ["disabled", disabled],
        ["duration (ms)", duration],
        ["exit code", rc],
    ]
    xlsx = write_excel(os.path.join(args.outdir, "report_offical.xlsx"),
                       meta_rows, cases)
    if xlsx:
        print(f"excel report : {xlsx}")
    else:
        print("(openpyxl unavailable; raw log kept at", log_path, ")")
    return rc


if __name__ == "__main__":
    sys.exit(main())
