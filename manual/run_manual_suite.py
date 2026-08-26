#!/usr/bin/env python3
"""Unified driver for the OpenCV camera MANUAL test suite.

Orchestrates the 6 applicable items from
opencv_claude/camera_api_test/manifest_manual.yaml:

  item                              automation level
  --------------------------------  ------------------------------------------------
  device_index_physical_mapping     AUTO-evidence (sysfs enumeration + node probing)
  exposure_visual_check             AUTO-evidence (sweep helper, brightness metric)
  autofocus_visual_check            AUTO-evidence (focus sweep, Laplacian sharpness)
  white_balance_visual_check        AUTO-evidence (WB temp sweep, RGB cast metric)
  real_usb_unplug_reconnect         ASSISTED   (runs reconnect tool inside a time
                                               window while YOU unplug/replug)
  long_duration_stability           OPTIONAL   (only with --long-run MINUTES)

Flow per item: run automated evidence collection -> show summary ->
operator enters verdict (p/f/s) -> merged into report/new/report_manual.json
(same schema as manual_test_runner.py, so report_generator.py consumes it).

Usage:
  python3 run_manual_suite.py --list
  python3 run_manual_suite.py                       # all items, interactive
  python3 run_manual_suite.py --item exposure_visual_check
"""

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone

import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
def _find_claude_dir():
    candidates = [
        os.environ.get("CAMERA_API_CLAUDE_DIR", ""),
        os.path.join(HERE, "..", "opencv_claude", "camera_api_test"),
        os.path.join(HERE, "..", "..", "opencv_claude", "camera_api_test"),
    ]
    for c in candidates:
        if c and os.path.isfile(os.path.join(c, "manifest_manual.yaml")):
            return os.path.normpath(os.path.abspath(c))
    return candidates[1]


CLAUDE_TEST_DIR = _find_claude_dir()
CLAUDE_MANIFEST = None  # resolved lazily in main()
HELPERS = {
    "exposure": os.path.join(HERE, "exposure_visual_check.py"),
    "autofocus": os.path.join(HERE, "autofocus_visual_check.py"),
    "white_balance": os.path.join(HERE, "white_balance_visual_check.py"),
    "reconnect": os.path.join(HERE, "usb_unplug_reconnect_test.py"),
}
ITEMS_ORDER = [
    "device_index_physical_mapping",
    "exposure_visual_check",
    "autofocus_visual_check",
    "white_balance_visual_check",
    "real_usb_unplug_reconnect",
    "long_duration_stability",
]
REMOVED_ITEMS = {"multi_camera_physical_sync", "csi_camera_real_hardware"}


def get_opencv_source():
    # Detect source: env override from run_camera_tests, else python binding
    src = os.environ.get("REPORT_OPENCV_SOURCE") or os.environ.get("CPP_OPENCV_SOURCE") or os.environ.get("OPENCV_SOURCE")
    if src:
        return src
    # python bindings are pip-installed, treat as apt-like system
    return "apt"


def get_opencv_version():
    try:
        import cv2
        return cv2.__version__
    except Exception:
        return "unknown"


def hr(title=""):
    print("\n" + "=" * 72)
    if title:
        print(title)
        print("=" * 72)


def run_helper(script, argv, timeout=900):
    proc = subprocess.run([sys.executable, script] + argv,
                          capture_output=True, text=True, timeout=timeout)
    out = proc.stdout
    summary = None
    for line in reversed(out.splitlines()):
        if line.startswith("[") and line.rstrip().endswith("]"):
            try:
                candidate = json.loads(line)
                if isinstance(candidate, list):
                    summary = candidate
                    break
            except json.JSONDecodeError:
                continue
    tail = "\n".join(out.strip().splitlines()[-25:])
    return proc.returncode, summary, tail


def ask_verdict(item, auto_note, forced=None):
    print(f"\nAuto-measured evidence (will be written to note):\n{auto_note}")
    if forced:
        mapped = {"p": "PASS", "f": "FAIL", "s": "SKIP"}[forced.lower()]
        print(f"[non-interactive] auto verdict = {mapped}")
        return mapped, auto_note
    while True:
        try:
            raw = input(f"\n[{item}] Your verdict? (p=PASS / f=FAIL / "
                        f"s=SKIP / Enter=SKIP): ").strip().lower()
        except EOFError:
            print("  [EOF -> treated as SKIP]")
            return "SKIP", auto_note
        if raw in ("p", "pass"):
            return "PASS", auto_note
        if raw in ("f", "fail"):
            return "FAIL", auto_note
        if raw in ("", "s", "skip"):
            return "SKIP", auto_note
        print("please enter p / f / s")


# ---------------- item runners: return (status, note) ----------------

def ev_device_mapping(a):
    lines = ["Device node mapping:"]
    candidates = []
    for path in sorted(glob.glob("/sys/class/video4linux/video*")):
        node = "/dev/" + os.path.basename(path)
        try:
            name = open(os.path.join(path, "name")).read().strip()
        except OSError:
            name = "?"
        import cv2
        cap = cv2.VideoCapture(node, cv2.CAP_V4L2)
        opened = cap.isOpened()
        ok_frame = False
        if opened:
            ok_frame, _ = cap.read()
        cap.release()
        kind = "CAPTURE" if (opened and ok_frame) else "not-capturable"
        candidates.append((node, name, kind))
        lines.append(f"  {node:<14} {name:<38} {kind}")
    usable = [c for c in candidates if c[2] == "CAPTURE"]
    lines.append(f"capture-capable nodes: {', '.join(c[0] for c in usable)}")
    note = "\n".join(lines) + f" | target={a.device}"
    hr("[1/6] device_index_physical_mapping")
    print(note)
    return ask_verdict("device_index_physical_mapping", note, a.answer)


def ev_exposure(a):
    hr("[2/6] exposure_visual_check")
    rc, summ, tail = run_helper(
        HELPERS["exposure"],
        ["-d", a.device, "sweep", "--values", a.exposure_values,
         "--outdir", os.path.abspath(a.evidence)])
    print(tail)
    note = (f"target={a.device}; sweep values={a.exposure_values}; "
            f"helper_exit={rc}; measured={summ}")
    return ask_verdict("exposure_visual_check", note, a.answer)


def ev_autofocus(a):
    hr("[3/6] autofocus_visual_check")
    rc1, s1, t1 = run_helper(
        HELPERS["autofocus"],
        ["-d", a.device, "focus", "--values", a.focus_values,
         "--outdir", os.path.abspath(a.evidence)])
    print(t1)
    rc2, s2, t2 = run_helper(
        HELPERS["autofocus"],
        ["-d", a.device, "zoom", "--values", "100,150,200",
         "--outdir", os.path.abspath(a.evidence)])
    print(t2)
    note = (f"target={a.device}; FOCUS sweep={s1}; ZOOM sweep={s2}; "
            f"PAN/TILT needs separate confirmation (get stuck at 0 = unsupported)")
    return ask_verdict("autofocus_visual_check", note, a.answer)


def ev_white_balance(a):
    hr("[4/6] white_balance_visual_check")
    rc, summ, tail = run_helper(
        HELPERS["white_balance"],
        ["-d", a.device, "temp", "--values", a.wb_values,
         "--outdir", os.path.abspath(a.evidence)])
    print(tail)
    note = f"target={a.device}; WB_TEMP sweep={summ}"
    return ask_verdict("white_balance_visual_check", note, a.answer)


def ev_usb_reconnect(a):
    hr("[5/6] real_usb_unplug_reconnect  (interactive window)")
    print(f"Next {a.reconnect_window} seconds:\n"
          f"  1. After stream_established >>> UNPLUG the USB cable of {a.device}\n"
          f"  2. After disconnect_detected wait 5-10s, then replug\n"
          f"  3. Done once you see recovered\n\n"
          f"--- live status ---")
    env = dict(os.environ, PYTHONUNBUFFERED="1")
    proc = subprocess.Popen(
        [sys.executable, "-u", HELPERS["reconnect"],
         "--device", a.device,
         "--max-duration", str(a.reconnect_window),
         "--outdir", os.path.join(a.evidence, "reconnect")],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
        text=True, bufsize=1, env=env)
    stage = 0
    for line in proc.stdout:
        line = line.rstrip()
        if not line:
            continue
        print(f"    | {line}", flush=True)
        if stage == 0 and "stream_established" in line:
            stage = 1
            print("\n>>> NOW UNPLUG the USB cable! <<<\n", flush=True)
        elif stage == 1 and "disconnect_detected" in line:
            stage = 2
            print("\n>>> Disconnect detected! Wait 5-10s then replug USB <<<\n",
                  flush=True)
        elif stage == 2 and "recovered" in line:
            stage = 3
            print("\n>>> Stream recovered! Thanks - waiting for timer... <<<\n", flush=True)
    proc.wait()
    sj = os.path.join(a.evidence, "reconnect", "summary.json")
    note = f"target={a.device}; window={a.reconnect_window}s"
    if os.path.exists(sj):
        with open(sj) as f:
            data = json.load(f)
        note += (f"; disconnects={json.dumps(data.get('disconnects', []))}; "
                 f"exceptions={data.get('exception_count')}")
    return ask_verdict("real_usb_unplug_reconnect", note, a.answer)


def ev_long_run(a):
    hr("[6/6] long_duration_stability")
    if not a.long_run:
        note = "--long-run minutes not given; skipping long-run observation"
        print(note)
        return "SKIP", note
    import cv2
    deadline = time.time() + a.long_run * 60
    samples, fails, peak_rss = [], 0, 0
    cap = cv2.VideoCapture(int(a.device) if str(a.device).isdigit() else a.device,
                           cv2.CAP_V4L2)
    last_log = 0
    while time.time() < deadline:
        ok, _ = cap.read()
        fails += (not ok)
        if time.time() - last_log >= 60:
            rss = int(open("/proc/self/status").read()
                      .split("VmRSS:")[1].split()[0])
            peak_rss = max(peak_rss, rss)
            samples.append({"t_min": round((time.time() - deadline +
                                             a.long_run * 60) / 60, 1),
                            "rss_kb": rss})
            print(f"  [{samples[-1]['t_min']}min] RSS={rss}KB fails={fails}")
            last_log = time.time()
    cap.release()
    note = f"dur={a.long_run}min, samples={json.dumps(samples)}, fails={fails}"
    print(note)
    return ask_verdict("long_duration_stability", note, a.answer)


RUNNERS = {
    "device_index_physical_mapping": ev_device_mapping,
    "exposure_visual_check": ev_exposure,
    "autofocus_visual_check": ev_autofocus,
    "white_balance_visual_check": ev_white_balance,
    "real_usb_unplug_reconnect": ev_usb_reconnect,
    "long_duration_stability": ev_long_run,
}



def write_excel(results, path, generated_at, opencv_version=None, device=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    illegal_re = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")

    def safe(v):
        if isinstance(v, str):
            cleaned = illegal_re.sub("?", v)
            return cleaned if cleaned.strip() else "(empty)"
        return v

    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"),
             "FAIL": PatternFill("solid", fgColor="FFC7CE"),
             "SKIP": PatternFill("solid", fgColor="D9D9D9"),
             "NOT_RUN": PatternFill("solid", fgColor="F2F2F2")}
    fonts = {"PASS": Font(color="006100"), "FAIL": Font(color="9C0006", bold=True),
             "SKIP": Font(color="404040"), "NOT_RUN": Font(color="808080")}
    thin = Border(*[Side(style="thin", color="BFBFBF")] * 4)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    counts = {}
    for r in results:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    # resolve opencv version lazily if not passed
    if opencv_version is None:
        opencv_version = get_opencv_version()
    try:
        import platform as _pl
        python_version = _pl.python_version()
        platform_str = _pl.platform()
    except Exception:
        python_version = "unknown"
        platform_str = "unknown"
    rows = [
        ["OpenCV Camera Manual Test Report", ""],
        ["", ""],
        ["generated_at", generated_at],
        ["opencv_version", opencv_version],
        ["python_version", python_version],
        ["platform", platform_str],
        ["device", device or "-"],
        ["total items", len(results)],
        ["PASS", counts.get("PASS", 0)], ["FAIL", counts.get("FAIL", 0)],
        ["SKIP", counts.get("SKIP", 0)], ["NOT_RUN", counts.get("NOT_RUN", 0)],
    ]
    for r in rows:
        ws.append([safe(x) for x in r])
    ws["A1"].font = Font(bold=True, size=14)
    for i in range(3, len(rows) + 1):
        a = ws.cell(row=i, column=1)
        if a.value and not str(a.value).startswith("=="):
            a.font = Font(bold=True)
            a.border = thin
            ws.cell(row=i, column=2).border = thin
    for col, w in (("A", 24), ("B", 60)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Results")
    ws2.append(["test_ref", "description", "Status", "note (evidence)", "checked_at"])
    for r in results:
        ws2.append([safe(r.get("test_ref", "?")),
                    safe(r.get("description", "")),
                    safe(r.get("status", "?")),
                    safe(r.get("note", "")),
                    safe(r.get("checked_at") or "-")])
    for c in ws2[1]:
        c.fill, c.font = hdr_fill, Font(bold=True, color="FFFFFF")
        c.border = thin
    st_col = 3
    for row in ws2.iter_rows(min_row=2):
        st = row[st_col - 1].value
        if st in fills:
            row[st_col - 1].fill, row[st_col - 1].font = fills[st], fonts[st]
        for c in row:
            c.border = thin
            c.alignment = Alignment(vertical="top", wrap_text=(c.column == 4))
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    for i, w in enumerate((34, 46, 10, 90, 26), 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def _find_manifest():
    """local ./manifest_manual.yaml wins; fall back to opencv_claude dir."""
    local = os.path.join(HERE, "manifest_manual.yaml")
    if os.path.isfile(local):
        return local
    return os.path.join(_find_claude_dir(), "manifest_manual.yaml")


def list_video_nodes():
    nodes = []
    for path in sorted(glob.glob("/sys/class/video4linux/video*")):
        node = "/dev/" + os.path.basename(path)
        try:
            name = open(os.path.join(path, "name")).read().strip()
        except OSError:
            name = "?"
        nodes.append((node, name))
    return nodes


def ensure_device_exists(device):
    """Abort with a clear error if the target camera node is absent."""
    dev = str(device).strip()
    if dev.startswith("/dev/"):
        node = dev                       # explicit path: trust as-is
    elif dev.isdigit():
        node = f"/dev/video{dev}"
    else:
        return                           # URLs / pipelines: skip check
    if os.path.exists(node):
        print(f"[device] using {node}")
        return
    print(f"[ERROR] target device '{device}' ({node}) does not exist.",
          file=sys.stderr)
    nodes = list_video_nodes()
    if nodes:
        print("Available video nodes:", file=sys.stderr)
        for n, nm in nodes:
            print(f"  {n:<14} {nm}", file=sys.stderr)
    else:
        print("No /dev/video* nodes found - is a camera plugged in?",
              file=sys.stderr)
    print("Plug in the camera, or pass -d/--device with an existing node.",
          file=sys.stderr)
    sys.exit(2)


def load_manifest():
    with open(CLAUDE_MANIFEST, encoding="utf-8") as f:
        entries = {e["test_ref"]: e for e in yaml.safe_load(f)}
    return entries


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--list", action="store_true")
    p.add_argument("--item", choices=ITEMS_ORDER)
    p.add_argument("-d", "--device", default="/dev/video0")
    p.add_argument("--exposure-values", default="50,100,200,400,800")
    p.add_argument("--focus-values", default="0,80,160,250")
    p.add_argument("--wb-values", default="2500,4000,5500,6500")
    p.add_argument("--reconnect-window", type=float, default=30)
    p.add_argument("--long-run", type=int, default=None, help="minutes")
    p.add_argument("--evidence", default="./manual_evidence")
    p.add_argument("--report", default=os.path.join("report", "new",
                                                "report_manual.json"))
    p.add_argument("--answer", choices=["p", "f", "s"], default=None,
                   help="non-interactive mode: same answer for every item")
    p.add_argument("--excel-dir", default=os.path.join("report", "new"),
                   help="Excel report output folder")
    p.add_argument("--no-excel", action="store_true",
                   help="skip the Excel report")
    a = p.parse_args()

    if a.list:
        for ref in ITEMS_ORDER:
            e = load_manifest().get(ref, {})
            print(f"[{ref}] {e.get('description','')}")
        return 0

    ensure_device_exists(a.device)

    os.makedirs(a.evidence, exist_ok=True)
    global CLAUDE_MANIFEST
    CLAUDE_MANIFEST = _find_manifest()
    print(f"manifest: {CLAUDE_MANIFEST}")
    manifest = load_manifest()
    order = [a.item] if a.item else ITEMS_ORDER

    existing = {}
    if os.path.exists(a.report):
        with open(a.report, encoding="utf-8") as f:
            existing = {r["test_ref"]: r for r in json.load(f).get("results", [])}

    def save_now(extra):
        merged = {ref: {"test_ref": ref, "status": "NOT_RUN", "note": "",
                        "checked_at": None} for ref in ITEMS_ORDER}
        merged.update({k: v for k, v in existing.items()
                       if k not in REMOVED_ITEMS})
        for r in extra:
            merged[r["test_ref"]] = r
        # collect environment metadata
        opencv_ver = get_opencv_version()
        opencv_src = get_opencv_source()
        try:
            import platform as _pl
            env = {
                "opencv_version": opencv_ver,
                "opencv_source": opencv_src,
                "python": _pl.python_version(),
                "platform": _pl.platform(),
                "machine": _pl.machine(),
            }
        except Exception:
            env = {"opencv_version": opencv_ver, "opencv_source": opencv_src}
        report = {"report_type": "manual",
                  "generated_at": datetime.now(timezone.utc).isoformat(),
                  "opencv_version": opencv_ver,
                  "opencv_source": opencv_src,
                  "environment": env,
                  "device": str(a.device),
                  "results": [merged[r] for r in ITEMS_ORDER]}
        os.makedirs(os.path.dirname(os.path.abspath(a.report)) or ".",
                    exist_ok=True)
        with open(a.report, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        if not a.no_excel:
            write_excel(list(merged.values()),
                        os.path.join(a.excel_dir, "report_manual.xlsx"),
                        report["generated_at"],
                        opencv_version=opencv_ver,
                        device=str(a.device))
        return report

    results = []
    interrupted = False
    for ref in order:
        entry = manifest.get(ref)
        if entry is None or ref not in RUNNERS:
            continue
        try:
            status, note = RUNNERS[ref](a)
        except KeyboardInterrupt:
            print("\n[interrupted by user] saving partial report...")
            status, note = "INTERRUPTED", "interrupted by user (Ctrl+C)"
            interrupted = True
        except Exception as e:
            status, note = "SKIP", f"runner error: {type(e).__name__}: {e}"
        results.append({
            "test_ref": ref,
            "description": entry.get("description", ""),
            "status": status,
            "note": note,
            "checked_at": datetime.now(timezone.utc).isoformat(),
        })
        save_now(results)
        if interrupted:
            break

    report = save_now(results)          # final authoritative save
    merged = {r["test_ref"]: r for r in report["results"]}
    os.makedirs(a.excel_dir, exist_ok=True)
    xlsx_path = None
    if not a.no_excel:
        xlsx_path = write_excel(
            list(merged.values()),
            os.path.join(a.excel_dir, "report_manual.xlsx"),
            report["generated_at"],
            opencv_version=report.get("opencv_version", get_opencv_version()),
            device=report.get("device", str(a.device)))
        if xlsx_path is None:
            print("(openpyxl not installed; skipping Excel report - auto-enabled after pip install openpyxl)")

    done = sum(1 for r in merged.values() if r["status"] != "NOT_RUN")
    passed = sum(1 for r in merged.values() if r["status"] == "PASS")
    failed = sum(1 for r in merged.values() if r["status"] == "FAIL")
    hr("MANUAL SUITE SUMMARY")
    print(f"  executed this run : {len(results)}")
    print(f"  total recorded    : {done}/{len(ITEMS_ORDER)}  (PASS={passed}, FAIL={failed})")
    print(f"  report            : {a.report}")
    print(f"  excel             : see {a.excel_dir}/report_manual.xlsx")
    print(f"  evidence          : {a.evidence}/")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
